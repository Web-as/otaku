#!/usr/bin/env python3
"""
Metadata Extractor - Extracts metadata from various file types.
"""

import os
import sys
from pathlib import Path
from typing import Dict, Optional, Any
from datetime import datetime
import json

try:
    from PIL import Image
    from PIL.ExifTags import TAGS, GPSTAGS
    PILLOW_AVAILABLE = True
except ImportError:
    PILLOW_AVAILABLE = False

try:
    from mutagen import File as MutagenFile
    from mutagen.id3 import ID3NoHeaderError
    MUTAGEN_AVAILABLE = True
except ImportError:
    MUTAGEN_AVAILABLE = False

try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class MetadataExtractor:
    """Extracts metadata from various file types."""
    
    def __init__(self):
        """Initialize the metadata extractor."""
        self.supported_formats = {
            'images': ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff', '.tif', 
                      '.webp', '.heic', '.raw', '.cr2', '.nef'],
            'audio': ['.mp3', '.wav', '.flac', '.aac', '.ogg', '.wma', '.m4a',
                     '.opus', '.amr', '.aiff', '.au'],
            'video': ['.mp4', '.avi', '.mkv', '.mov', '.wmv', '.flv', '.webm',
                     '.m4v', '.mpg', '.mpeg', '.3gp', '.vob', '.ts', '.m2ts'],
            'pdf': ['.pdf'],
            'docx': ['.docx'],
            'xlsx': ['.xlsx', '.xls'],
            'pptx': ['.pptx', '.ppt'],
            'text': ['.txt', '.md', '.log', '.csv']
        }
    
    def extract(self, file_path: Path) -> Dict[str, Any]:
        """
        Extract metadata from a file.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Dictionary containing extracted metadata
        """
        metadata = {
            'file_path': str(file_path),
            'file_name': file_path.name,
            'file_extension': file_path.suffix.lower(),
            'file_size': 0,
            'creation_date': None,
            'modification_date': None,
            'author': None,
            'title': None,
            'genre': None,
            'date': None,
            'album': None,
            'subject': None,
            'keywords': None,
            'description': None,
            'duration': None,
            'resolution': None,
            'codec': None,
            'bitrate': None,
            'sample_rate': None,
            'channels': None,
            'camera': None,
            'location': None,
            'language': None,
            'encoding': None,
            'line_count': None,
        }
        
        try:
            # Get file system metadata
            stat = file_path.stat()
            metadata['file_size'] = stat.st_size
            metadata['creation_date'] = datetime.fromtimestamp(stat.st_ctime).isoformat()
            metadata['modification_date'] = datetime.fromtimestamp(stat.st_mtime).isoformat()
        except (OSError, PermissionError):
            pass
        
        # Extract type-specific metadata
        ext = file_path.suffix.lower()
        
        if ext in self.supported_formats['images']:
            metadata.update(self._extract_image_metadata(file_path))
        elif ext in self.supported_formats['audio']:
            metadata.update(self._extract_audio_metadata(file_path))
        elif ext in self.supported_formats['video']:
            metadata.update(self._extract_video_metadata(file_path))
        elif ext in self.supported_formats['pdf']:
            metadata.update(self._extract_pdf_metadata(file_path))
        elif ext in self.supported_formats['docx']:
            metadata.update(self._extract_docx_metadata(file_path))
        elif ext in self.supported_formats['xlsx']:
            metadata.update(self._extract_xlsx_metadata(file_path))
        elif ext in self.supported_formats['pptx']:
            metadata.update(self._extract_pptx_metadata(file_path))
        elif ext in self.supported_formats['text']:
            metadata.update(self._extract_text_metadata(file_path))
        
        return metadata
    
    def _extract_image_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from image files."""
        metadata = {}
        
        if not PILLOW_AVAILABLE:
            return metadata
        
        try:
            with Image.open(file_path) as img:
                # Basic image info
                metadata['resolution'] = f"{img.width}x{img.height}"
                metadata['format'] = img.format
                metadata['mode'] = img.mode
                
                # EXIF data
                exif_data = img._getexif()
                if exif_data:
                    for tag_id, value in exif_data.items():
                        tag = TAGS.get(tag_id, tag_id)
                        
                        if tag == 'DateTime':
                            metadata['date'] = str(value)
                        elif tag == 'DateTimeOriginal':
                            metadata['date'] = str(value)
                        elif tag == 'Artist':
                            metadata['author'] = str(value)
                        elif tag == 'Make' or tag == 'Model':
                            camera_info = []
                            if 'Make' in [TAGS.get(t, t) for t in exif_data.keys()]:
                                make = exif_data.get([k for k, v in TAGS.items() if v == 'Make'][0])
                                if make:
                                    camera_info.append(str(make))
                            if 'Model' in [TAGS.get(t, t) for t in exif_data.keys()]:
                                model = exif_data.get([k for k, v in TAGS.items() if v == 'Model'][0])
                                if model:
                                    camera_info.append(str(model))
                            if camera_info:
                                metadata['camera'] = ' '.join(camera_info)
                        elif tag == 'GPSInfo':
                            # Extract GPS coordinates if available
                            gps_data = {}
                            for gps_tag_id, gps_value in value.items():
                                gps_tag = GPSTAGS.get(gps_tag_id, gps_tag_id)
                                gps_data[gps_tag] = gps_value
                            if gps_data:
                                metadata['location'] = json.dumps(gps_data)
        except Exception as e:
            # Silently fail - return empty metadata
            pass
        
        return metadata
    
    def _extract_audio_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from audio files."""
        metadata = {}
        
        if not MUTAGEN_AVAILABLE:
            return metadata
        
        try:
            audio_file = MutagenFile(file_path)
            if audio_file is None:
                return metadata
            
            # Common tags
            if hasattr(audio_file, 'tags'):
                tags = audio_file.tags
                
                # Try common tag names
                tag_mapping = {
                    'artist': ['TPE1', 'ARTIST', '©ART'],
                    'album': ['TALB', 'ALBUM', '©alb'],
                    'title': ['TIT2', 'TITLE', '©nam'],
                    'genre': ['TCON', 'GENRE', '©gen'],
                    'date': ['TDRC', 'DATE', 'YEAR', '©day'],
                    'track': ['TRCK', 'TRACK', 'TRACKNUMBER'],
                }
                
                for key, tag_names in tag_mapping.items():
                    for tag_name in tag_names:
                        if tag_name in tags:
                            value = tags[tag_name]
                            if isinstance(value, list) and len(value) > 0:
                                value = value[0]
                            if value:
                                if key == 'artist':
                                    metadata['author'] = str(value)
                                elif key == 'date':
                                    metadata['date'] = str(value)
                                else:
                                    metadata[key] = str(value)
                                break
            
            # Audio properties
            if hasattr(audio_file, 'info'):
                info = audio_file.info
                if hasattr(info, 'length'):
                    metadata['duration'] = info.length
                if hasattr(info, 'bitrate'):
                    metadata['bitrate'] = info.bitrate
                if hasattr(info, 'sample_rate'):
                    metadata['sample_rate'] = info.sample_rate
                if hasattr(info, 'channels'):
                    metadata['channels'] = info.channels
        except (ID3NoHeaderError, Exception):
            # Silently fail
            pass
        
        return metadata
    
    def _extract_video_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from video files."""
        metadata = {}
        
        if not MUTAGEN_AVAILABLE:
            return metadata
        
        try:
            video_file = MutagenFile(file_path)
            if video_file is None:
                return metadata
            
            # Video tags
            if hasattr(video_file, 'tags'):
                tags = video_file.tags
                
                tag_mapping = {
                    'title': ['TITL', 'TITLE', '©nam'],
                    'artist': ['ARTI', 'ARTIST', '©ART'],
                    'genre': ['GNRE', 'GENRE', '©gen'],
                    'date': ['DATE', 'YEAR', '©day'],
                }
                
                for key, tag_names in tag_mapping.items():
                    for tag_name in tag_names:
                        if tag_name in tags:
                            value = tags[tag_name]
                            if isinstance(value, list) and len(value) > 0:
                                value = value[0]
                            if value:
                                if key == 'artist':
                                    metadata['author'] = str(value)
                                elif key == 'date':
                                    metadata['date'] = str(value)
                                else:
                                    metadata[key] = str(value)
                                break
            
            # Video properties
            if hasattr(video_file, 'info'):
                info = video_file.info
                if hasattr(info, 'length'):
                    metadata['duration'] = info.length
                if hasattr(info, 'bitrate'):
                    metadata['bitrate'] = info.bitrate
                if hasattr(info, 'width') and hasattr(info, 'height'):
                    metadata['resolution'] = f"{info.width}x{info.height}"
        except Exception:
            # Silently fail
            pass
        
        return metadata
    
    def _extract_pdf_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from PDF files."""
        metadata = {}
        
        if not PYPDF2_AVAILABLE:
            return metadata
        
        try:
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                
                if pdf_reader.metadata:
                    meta = pdf_reader.metadata
                    
                    if '/Author' in meta:
                        metadata['author'] = str(meta['/Author'])
                    if '/Title' in meta:
                        metadata['title'] = str(meta['/Title'])
                    if '/Subject' in meta:
                        metadata['subject'] = str(meta['/Subject'])
                    if '/Keywords' in meta:
                        metadata['keywords'] = str(meta['/Keywords'])
                    if '/CreationDate' in meta:
                        metadata['date'] = str(meta['/CreationDate'])
                    if '/ModDate' in meta:
                        if not metadata.get('date'):
                            metadata['date'] = str(meta['/ModDate'])
        except Exception:
            # Silently fail
            pass
        
        return metadata
    
    def _extract_docx_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from Word documents."""
        metadata = {}
        
        if not DOCX_AVAILABLE:
            return metadata
        
        try:
            doc = DocxDocument(file_path)
            core_props = doc.core_properties
            
            if core_props.author:
                metadata['author'] = core_props.author
            if core_props.title:
                metadata['title'] = core_props.title
            if core_props.subject:
                metadata['subject'] = core_props.subject
            if core_props.keywords:
                metadata['keywords'] = core_props.keywords
            if core_props.created:
                metadata['date'] = core_props.created.isoformat()
            if core_props.modified:
                if not metadata.get('date'):
                    metadata['date'] = core_props.modified.isoformat()
        except Exception:
            # Silently fail
            pass
        
        return metadata
    
    def _extract_xlsx_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from Excel files."""
        metadata = {}
        
        if not OPENPYXL_AVAILABLE:
            return metadata
        
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            props = workbook.properties
            
            if props.creator:
                metadata['author'] = props.creator
            if props.title:
                metadata['title'] = props.title
            if props.subject:
                metadata['subject'] = props.subject
            if props.keywords:
                metadata['keywords'] = props.keywords
            if props.created:
                metadata['date'] = props.created.isoformat()
            if props.modified:
                if not metadata.get('date'):
                    metadata['date'] = props.modified.isoformat()
        except Exception:
            # Silently fail
            pass
        
        return metadata
    
    def _extract_pptx_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata from PowerPoint files."""
        metadata = {}
        
        if not PPTX_AVAILABLE:
            return metadata
        
        try:
            prs = Presentation(file_path)
            core_props = prs.core_properties
            
            if core_props.author:
                metadata['author'] = core_props.author
            if core_props.title:
                metadata['title'] = core_props.title
            if core_props.subject:
                metadata['subject'] = core_props.subject
            if core_props.keywords:
                metadata['keywords'] = core_props.keywords
            if core_props.created:
                metadata['date'] = core_props.created.isoformat()
            if core_props.modified:
                if not metadata.get('date'):
                    metadata['date'] = core_props.modified.isoformat()
        except Exception:
            # Silently fail
            pass
        
        return metadata
    
    def _extract_text_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract basic metadata from text files."""
        metadata = {}
        
        try:
            # Try to detect encoding and count lines
            encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        lines = f.readlines()
                        metadata['encoding'] = encoding
                        metadata['line_count'] = len(lines)
                        break
                except (UnicodeDecodeError, UnicodeError):
                    continue
            
            # Basic language detection (simple heuristic)
            if metadata.get('line_count', 0) > 0:
                try:
                    with open(file_path, 'r', encoding=metadata.get('encoding', 'utf-8')) as f:
                        sample = f.read(1000)
                        # Simple heuristic: check for common words
                        if any(word in sample.lower() for word in ['the', 'and', 'is', 'are']):
                            metadata['language'] = 'en'
                except:
                    pass
        except Exception:
            # Silently fail
            pass
        
        return metadata
    
    def get_standardized_metadata(self, file_path: Path) -> Dict[str, Any]:
        """
        Get standardized metadata dictionary.
        
        Args:
            file_path: Path to the file
            
        Returns:
            Standardized metadata dictionary with common fields
        """
        raw_metadata = self.extract(file_path)
        
        # Standardize field names
        standardized = {
            'author': raw_metadata.get('author') or raw_metadata.get('artist') or 'Unknown',
            'genre': raw_metadata.get('genre') or 'Unknown',
            'date': raw_metadata.get('date') or raw_metadata.get('creation_date') or 'Unknown',
            'title': raw_metadata.get('title') or raw_metadata.get('file_name', 'Unknown'),
            'album': raw_metadata.get('album') or 'Unknown',
            'subject': raw_metadata.get('subject') or 'Unknown',
            'description': raw_metadata.get('description') or raw_metadata.get('keywords') or 'Unknown',
            'duration': raw_metadata.get('duration'),
            'resolution': raw_metadata.get('resolution'),
            'codec': raw_metadata.get('codec'),
            'bitrate': raw_metadata.get('bitrate'),
            'sample_rate': raw_metadata.get('sample_rate'),
            'channels': raw_metadata.get('channels'),
            'camera': raw_metadata.get('camera'),
            'location': raw_metadata.get('location'),
            'file_size': raw_metadata.get('file_size', 0),
            'file_path': raw_metadata.get('file_path'),
            'file_name': raw_metadata.get('file_name'),
        }
        
        # Clean up "Unknown" values - convert to None for better handling
        for key, value in standardized.items():
            if value == 'Unknown':
                standardized[key] = None
        
        return standardized

